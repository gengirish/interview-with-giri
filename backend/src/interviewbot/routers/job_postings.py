import contextlib
import csv
from datetime import datetime
import io
import secrets
from uuid import UUID

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, status
from pydantic import BaseModel, ValidationError
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from interviewbot.config import get_settings
from interviewbot.dependencies import get_db, get_org_id, require_role
from interviewbot.models.schemas import (
    InterviewConfig,
    InterviewFormat,
    JobPostingCreateRequest,
    JobPostingResponse,
    JobPostingUpdateRequest,
    PaginatedResponse,
    RoleType,
)
from interviewbot.models.tables import InterviewSession, JobPosting, Organization

router = APIRouter(prefix="/job-postings", tags=["Job Postings"])


class GenerateLinkRequest(BaseModel):
    candidate_name: str | None = None
    candidate_email: str | None = None
    scheduled_at: datetime | None = None


IMPORT_COLUMNS = [
    "title",
    "role_type",
    "job_description",
    "required_skills",
    "interview_format",
    "num_questions",
    "duration_minutes",
    "difficulty",
    "include_coding",
]


@router.post("", response_model=JobPostingResponse, status_code=status.HTTP_201_CREATED)
async def create_job_posting(
    req: JobPostingCreateRequest,
    user: dict = Depends(require_role("admin", "hiring_manager")),
    db: AsyncSession = Depends(get_db),
    org_id: UUID = Depends(get_org_id),
) -> JobPostingResponse:
    posting = JobPosting(
        org_id=org_id,
        title=req.title,
        role_type=req.role_type.value,
        job_description=req.job_description,
        required_skills=req.required_skills,
        interview_format=req.interview_format.value,
        interview_config=req.interview_config.model_dump(),
        scoring_rubric=req.scoring_rubric,
    )
    db.add(posting)
    await db.commit()
    await db.refresh(posting)

    return _to_response(posting)


@router.get("", response_model=PaginatedResponse)
async def list_job_postings(
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1, le=100),
    q: str | None = Query(None),
    is_active: bool | None = Query(None),
    role_type: str | None = Query(None),
    interview_format: str | None = Query(None),
    user: dict = Depends(require_role("admin", "hiring_manager", "viewer")),
    db: AsyncSession = Depends(get_db),
    org_id: UUID = Depends(get_org_id),
) -> PaginatedResponse:
    base_query = select(JobPosting).where(JobPosting.org_id == org_id)

    if q:
        base_query = base_query.where(
            or_(
                JobPosting.title.ilike(f"%{q}%"),
                JobPosting.job_description.ilike(f"%{q}%"),
            )
        )
    if is_active is not None:
        base_query = base_query.where(JobPosting.is_active == is_active)
    if role_type:
        base_query = base_query.where(JobPosting.role_type == role_type)
    if interview_format:
        base_query = base_query.where(JobPosting.interview_format == interview_format)

    count_result = await db.execute(select(func.count()).select_from(base_query.subquery()))
    total = count_result.scalar() or 0

    result = await db.execute(
        base_query.order_by(JobPosting.created_at.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
    )
    postings = result.scalars().all()

    return PaginatedResponse(
        items=[_to_response(p) for p in postings],
        total=total,
        page=page,
        per_page=per_page,
    )


@router.get("/{posting_id}", response_model=JobPostingResponse)
async def get_job_posting(
    posting_id: UUID,
    user: dict = Depends(require_role("admin", "hiring_manager", "viewer")),
    db: AsyncSession = Depends(get_db),
    org_id: UUID = Depends(get_org_id),
) -> JobPostingResponse:
    posting = await _get_posting_or_404(db, org_id, posting_id)
    return _to_response(posting)


@router.patch("/{posting_id}", response_model=JobPostingResponse)
async def update_job_posting(
    posting_id: UUID,
    req: JobPostingUpdateRequest,
    user: dict = Depends(require_role("admin", "hiring_manager")),
    db: AsyncSession = Depends(get_db),
    org_id: UUID = Depends(get_org_id),
) -> JobPostingResponse:
    posting = await _get_posting_or_404(db, org_id, posting_id)

    update_data = req.model_dump(exclude_unset=True)
    for field, value in update_data.items():
        if field == "interview_config" and value is not None:
            value = value.model_dump() if hasattr(value, "model_dump") else value
        if field == "interview_format" and value is not None:
            value = value.value if hasattr(value, "value") else value
        setattr(posting, field, value)

    await db.commit()
    await db.refresh(posting)
    return _to_response(posting)


@router.delete("/{posting_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_job_posting(
    posting_id: UUID,
    user: dict = Depends(require_role("admin")),
    db: AsyncSession = Depends(get_db),
    org_id: UUID = Depends(get_org_id),
) -> None:
    posting = await _get_posting_or_404(db, org_id, posting_id)
    await db.delete(posting)
    await db.commit()


@router.post("/{posting_id}/generate-link", response_model=dict)
async def generate_interview_link(
    posting_id: UUID,
    body: GenerateLinkRequest | None = None,
    user: dict = Depends(require_role("admin", "hiring_manager")),
    db: AsyncSession = Depends(get_db),
    org_id: UUID = Depends(get_org_id),
) -> dict:
    posting = await _get_posting_or_404(db, org_id, posting_id)

    token = secrets.token_urlsafe(32)
    session = InterviewSession(
        job_posting_id=posting_id,
        org_id=org_id,
        token=token,
        candidate_name=body.candidate_name if body else None,
        candidate_email=body.candidate_email if body else None,
        scheduled_at=body.scheduled_at if body else None,
    )
    db.add(session)
    await db.commit()

    result: dict = {"token": token, "interview_url": f"/interview/{token}"}

    # Add scheduled_at whenever set; generate ICS and send invite when candidate_email present
    if body and body.scheduled_at:
        result["scheduled_at"] = body.scheduled_at.isoformat()
        if body.candidate_email:
            from interviewbot.services.calendar_service import generate_ics_invite

            config = posting.interview_config or {}
            duration = config.get("duration_minutes", 30)
            settings = get_settings()
            interview_url = f"{settings.app_url}/interview/{token}"

            ics_content = generate_ics_invite(
                summary=f"Interview: {posting.title}",
                description=f"AI Interview for {posting.title}\nJoin: {interview_url}",
                start_time=body.scheduled_at,
                duration_minutes=duration,
                attendee_email=body.candidate_email,
                location=interview_url,
            )
            result["ics_content"] = ics_content

            with contextlib.suppress(Exception):
                from interviewbot.services.notifications import send_interview_invitation

                org_result = await db.execute(
                    select(Organization).where(Organization.id == org_id)
                )
                org = org_result.scalar_one_or_none()
                org_name = org.name if org else "Company"
                org_inbox_id = org.agentmail_inbox_id if org else None

                await send_interview_invitation(
                    body.candidate_email,
                    body.candidate_name or "Candidate",
                    posting.title,
                    interview_url,
                    org_name=org_name,
                    org_inbox_id=org_inbox_id,
                )

    return result


@router.post("/{posting_id}/extract-skills", response_model=dict)
async def extract_skills(
    posting_id: UUID,
    user: dict = Depends(require_role("admin", "hiring_manager")),
    db: AsyncSession = Depends(get_db),
    org_id: UUID = Depends(get_org_id),
) -> dict:
    posting = await _get_posting_or_404(db, org_id, posting_id)

    from interviewbot.services.ai_engine import SKILL_EXTRACTION_PROMPT, AIEngine

    engine = AIEngine()
    prompt = SKILL_EXTRACTION_PROMPT.format(job_description=posting.job_description[:3000])

    import json as json_mod

    raw = await engine.chat([{"role": "user", "content": prompt}], temperature=0.3)

    try:
        cleaned = raw.strip()
        if cleaned.startswith("```"):
            cleaned = cleaned.split("\n", 1)[1].rsplit("```", 1)[0]
        result = json_mod.loads(cleaned)
    except json_mod.JSONDecodeError:
        result = {
            "technical_skills": [],
            "soft_skills": [],
            "error": "Failed to parse AI response",
        }

    all_skills = result.get("technical_skills", []) + result.get("soft_skills", [])
    if all_skills:
        posting.required_skills = all_skills
        await db.commit()

    return result


@router.post("/import", status_code=status.HTTP_200_OK)
async def import_job_postings(
    file: UploadFile,
    user: dict = Depends(require_role("admin", "hiring_manager")),
    db: AsyncSession = Depends(get_db),
    org_id: UUID = Depends(get_org_id),
) -> dict:
    """Bulk-import job postings from a CSV or Excel (.xlsx) file."""
    if not file.filename:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No file provided")

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ("csv", "xlsx"):
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Unsupported file format. Please upload a .csv or .xlsx file.",
        )

    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "File too large (max 5 MB)")

    rows: list[dict[str, str]] = []
    try:
        if ext == "csv":
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            rows = list(reader)
        else:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            if ws is None:
                raise HTTPException(status.HTTP_400_BAD_REQUEST, "Excel file has no active sheet")
            header: list[str] = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    header = [str(c or "").strip().lower().replace(" ", "_") for c in row]
                    continue
                row_dict = {header[j]: str(v or "") for j, v in enumerate(row) if j < len(header)}
                if any(v.strip() for v in row_dict.values()):
                    rows.append(row_dict)
            wb.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Failed to parse file: {exc}") from exc

    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "File contains no data rows")
    if len(rows) > 200:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Maximum 200 rows per import")

    results: list[dict] = []
    created_count = 0

    for idx, raw in enumerate(rows, start=2):
        row_num = idx
        title = raw.get("title", "").strip()
        if not title:
            results.append({"row": row_num, "status": "error", "error": "Missing title"})
            continue

        role_type_raw = raw.get("role_type", "mixed").strip().lower()
        job_desc = raw.get("job_description", "").strip()
        skills_raw = raw.get("required_skills", "").strip()
        fmt_raw = raw.get("interview_format", "text").strip().lower()
        num_q = raw.get("num_questions", "10").strip()
        dur = raw.get("duration_minutes", "30").strip()
        diff = raw.get("difficulty", "medium").strip().lower()
        coding_raw = raw.get("include_coding", "false").strip().lower()

        try:
            req = JobPostingCreateRequest(
                title=title,
                role_type=RoleType(role_type_raw) if role_type_raw else RoleType.MIXED,
                job_description=job_desc,
                required_skills=[s.strip() for s in skills_raw.split(",") if s.strip()]
                if skills_raw
                else [],
                interview_format=InterviewFormat(fmt_raw) if fmt_raw else InterviewFormat.TEXT,
                interview_config=InterviewConfig(
                    num_questions=int(num_q) if num_q else 10,
                    duration_minutes=int(dur) if dur else 30,
                    difficulty=diff or "medium",
                    include_coding=coding_raw in ("true", "1", "yes"),
                ),
            )
        except (ValidationError, ValueError) as e:
            msg = str(e.errors()[0]["msg"]) if isinstance(e, ValidationError) else str(e)
            results.append({"row": row_num, "title": title, "status": "error", "error": msg})
            continue

        posting = JobPosting(
            org_id=org_id,
            title=req.title,
            role_type=req.role_type.value,
            job_description=req.job_description,
            required_skills=req.required_skills,
            interview_format=req.interview_format.value,
            interview_config=req.interview_config.model_dump(),
        )
        db.add(posting)
        created_count += 1
        results.append({"row": row_num, "title": title, "status": "created"})

    if created_count > 0:
        await db.commit()

    return {
        "total_rows": len(rows),
        "created": created_count,
        "errors": len(rows) - created_count,
        "results": results,
    }


@router.get("/import/template")
async def download_import_template(
    user: dict = Depends(require_role("admin", "hiring_manager")),
) -> dict:
    """Return the expected CSV column headers for bulk import."""
    return {
        "columns": IMPORT_COLUMNS,
        "sample_row": {
            "title": "Senior Backend Engineer",
            "role_type": "technical",
            "job_description": (
                "We are looking for a senior backend engineer"
                " with 5+ years of experience in Python, FastAPI, and PostgreSQL."
            ),
            "required_skills": "Python, FastAPI, PostgreSQL, Docker",
            "interview_format": "text",
            "num_questions": "10",
            "duration_minutes": "30",
            "difficulty": "medium",
            "include_coding": "false",
        },
    }


async def _get_posting_or_404(db: AsyncSession, org_id: UUID, posting_id: UUID) -> JobPosting:
    result = await db.execute(
        select(JobPosting).where(
            JobPosting.id == posting_id,
            JobPosting.org_id == org_id,
        )
    )
    posting = result.scalar_one_or_none()
    if not posting:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Job posting not found")
    return posting


def _to_response(posting: JobPosting) -> JobPostingResponse:
    return JobPostingResponse(
        id=posting.id,
        org_id=posting.org_id,
        title=posting.title,
        role_type=posting.role_type,
        job_description=posting.job_description,
        required_skills=posting.required_skills or [],
        interview_format=posting.interview_format,
        interview_config=posting.interview_config or {},
        scoring_rubric=posting.scoring_rubric,
        is_active=posting.is_active,
        created_at=posting.created_at,
    )
